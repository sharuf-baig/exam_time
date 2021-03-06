#importing all the required libraries
from flask import Flask, session,request,render_template,redirect,flash,url_for,send_file
from flask_session import Session
from sqlalchemy import create_engine
from sqlalchemy.orm import scoped_session, sessionmaker
import os
from models import *
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.utils import secure_filename
import requests
import time
from functools import wraps
from wtforms import Form, StringField, TextAreaField, PasswordField, validators, DateTimeField, BooleanField, IntegerField
from flask_wtf import FlaskForm
from flask_wtf.file import FileField, FileRequired, FileAllowed
from wtforms_components import TimeField
from wtforms.fields.html5 import DateField
from docx import Document
from wtforms.validators import ValidationError
from coolname import generate_slug
from datetime import timedelta, datetime
import json
import csv
import random
from sqlalchemy import distinct
from sqlalchemy.sql import func
engine = create_engine(os.getenv("DATABASE_URL")) # database engine object from SQLAlchemy that manages connections to the database                                                  # DATABASE_URL is an environment variable that indicates where the database lives
db2 = scoped_session(sessionmaker(bind=engine))
#####################################
app = Flask(__name__)
app.secret_key= 'temtemp'
# Check for environment variable
if not os.getenv("DATABASE_URL"):
    raise RuntimeError("DATABASE_URL is not set")

# Configure session to use filesystem
app.config["SESSION_PERMANENT"] = False
app.config["SESSION_TYPE"] = "filesystem"
Session(app)
# Tell Flask what SQLAlchemy database to use.
app.config["SQLALCHEMY_DATABASE_URI"] = os.getenv("DATABASE_URL")
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False
# Link the Flask app with the database (no Flask app is actually being run yet).
db.init_app(app)
#vefify for login
def is_logged(f):
	@wraps(f)
	def wrap(*args, **kwargs):
		if 'user' in session is not None:
			return f(*args, **kwargs)
		else:
			flash('Please login')
			return redirect('/')
	return wrap
def is_admin(f):
	@wraps(f)
	def wrap(*args, **kwargs):
		result = Admin.query.filter_by(name=session['user']).first()
		if result is not None:
			return f(*args, **kwargs)
		else:
			flash('Please check with admin')
			return redirect('/')
	return wrap
def doctodict(filepath):
	document = Document(filepath)
	data={}
	count=1
	for table in document.tables:
		temp = {}
		for rowNo,_ in enumerate(table.rows):
			temp[table.cell(rowNo, 0).text]=table.cell(rowNo, 1).text
		data[count] = temp
		count+=1
 
	return data

class UploadForm(FlaskForm):
	subject = StringField('Subject')
	topic = StringField('Topic')
	doc = FileField('Docx Upload', validators=[FileRequired()])
	start_date = DateField('Start Date')
	start_time = TimeField('Start Time', default=datetime.utcnow()+timedelta(hours=5.5))
	end_date = DateField('End Date')
	end_time = TimeField('End Time', default=datetime.utcnow()+timedelta(hours=5.5))
	show_result = BooleanField('Show Result after completion')
	neg_mark = BooleanField('Enable negative marking')
	duration = IntegerField('Duration(in min)')
	password = StringField('Test Password', [validators.Length(min=3, max=6)])

	def validate_end_date(form, field):
		if field.data < form.start_date.data:
			raise ValidationError("End date must not be earlier than start date.")
	
	def validate_end_time(form, field):
		start_date_time = datetime.strptime(str(form.start_date.data) + " " + str(form.start_time.data),"%Y-%m-%d %H:%M:%S").strftime("%Y-%m-%d %H:%M")
		end_date_time = datetime.strptime(str(form.end_date.data) + " " + str(field.data),"%Y-%m-%d %H:%M:%S").strftime("%Y-%m-%d %H:%M")
		if start_date_time >= end_date_time:
			raise ValidationError("End date time must not be earlier/equal than start date time")
	
	def validate_start_date(form, field):
		if datetime.strptime(str(form.start_date.data) + " " + str(form.start_time.data),"%Y-%m-%d %H:%M:%S") < datetime.now():
			raise ValidationError("Start date and time must not be earlier than current")

class TestForm(Form):
	password = PasswordField('Test Password')

#various routes for the website
#login part
@app.route("/",methods=["GET","POST"])
def login():
	if session.get("user") is None:
		if request.method == "GET":
			return render_template('index.html',info=['login','UserName','Password','register'])
		else:
			name = request.form.get("name")
			password =request.form.get("password")
			user = User.query.filter_by(name=name).first()
			if user is None:
				flash('No user found with that sort of name')
				return redirect("/")
			elif not check_password_hash(user.password,password):
				flash("you password is incorrect")
				return redirect("/")
			else:
				session["user"]=user.name
				return redirect('/home')
	else:
		return redirect('/home')
@app.route("/register",methods=["GET","POST"])		
def register():
	if session.get("user") is None:
		if request.method == "GET":
			return render_template('index.html',info=['register','Your UserName','Your Password','login'])
		elif request.method == "POST":
			name = request.form.get("name")
			password =request.form.get("password")
			email =request.form.get("email")
			password_hash = generate_password_hash(password)
			user = User.query.filter_by(name=name).first() or User.query.filter_by(email=email).first()
			if user is None: 
				user = User(name=name,password=password_hash,email=email)
				session["user"]=user.name	
				db.session.add(user)
				db.session.commit()
				return redirect('/home')
			else:
				flash('User Name/Email is Taken')
				return redirect('/register')	
	else:
		return redirect('/home')
@app.route("/home",methods=["GET","POST"])
@is_logged
def home():
	return render_template('dashboard.html',info=['show_test','create_test','logout'])            
@app.route("/logout")
@is_logged
def logout():
	session.clear()
	return redirect("/")
@app.route('/create-test', methods = ['GET', 'POST'])
@is_logged
@is_admin
def create_test():
	form = UploadForm()
	if request.method == 'POST' and form.validate_on_submit():
		f = form.doc.data
		filename = secure_filename(f.filename)
		f.save('questions/' + filename)
		d = doctodict('questions/' + f.filename.replace(' ', '_').replace('(','').replace(')',''))
		test_id = generate_slug(2)
		try:
			for no, data in d.items():
				marks = data['((MARKS)) (1/2/3...)']
				a = data['((OPTION_A))']
				b = data['((OPTION_B))']
				c = data['((OPTION_C))']
				d = data['((OPTION_D))']
				question = data['((QUESTION))']
				correct_ans = data['((CORRECT_CHOICE)) (A/B/C/D)']
				explanation = data['((EXPLANATION)) (OPTIONAL)']

				tquestion = T_question(test_id=test_id,question_id=no,question=question,a=a,b=b,c=c,d=d,c_ans=correct_ans,marks=marks,explanation=explanation)
				db.session.add(tquestion)
				db.session.commit()
			start_date = form.start_date.data
			end_date = form.end_date.data
			start_time = form.start_time.data
			end_time = form.end_time.data
			start_date_time = str(start_date) + " " + str(start_time)
			end_date_time = str(end_date) + " " + str(end_time)
			show_result = form.show_result.data
			neg_mark = form.neg_mark.data

			duration = int(form.duration.data)*60
			password = form.password.data
			subject = form.subject.data
			topic = form.topic.data
			test=Test(test_id=test_id,name=session['user'],start=start_date_time,end=end_date_time,duration=duration,show_result=show_result,password=password,subject=subject,topic=topic,neg=neg_mark)
			db.session.add(test)
			db.session.commit()
			flash(f'Test ID: {test_id}', 'success')
			return redirect(url_for('home'))
		except Exception as e:
			flash('Invalid Input File Format','danger')
			return redirect(url_for('create_test'))
		
	return render_template('create_test.html' , form = form)
@app.route('/show-test', methods = ['GET', 'POST'])
@is_logged
def show_test():
	tests = Test.query.all()
	return render_template('show_tests.html',tests=tests)

@app.route("/give-test", methods = ['GET', 'POST'])
@is_logged
def give_test():
	global duration, marked_ans	
	name=session['user']
	form = TestForm(request.form)
	val = request.args.get('val', None)
	results = Test.query.filter_by(test_id=val).first()
	new=results
	duration1=new.duration
	start1=str(new.start)
	end1=str(new.end)
	neg1=new.neg
	data1={'duration': duration1,'negmarks':neg1,'start':start1,'end':end1}
	
	if request.method == 'POST' and form.validate():
		test_id = val
		
		print(results)
		if test_id is None:
			return redirect(url_for('show_test'))
		password_candidate = form.password.data
		results = Test.query.filter_by(test_id=test_id).first()
		if results is not None:
			data = results
			password = data.password
			duration = data.duration
			start = data.start
			start = str(start)
			end = data.end
			end = str(end)
			if password == password_candidate:
				now = datetime.now()
				now = now.strftime("%Y-%m-%d %H:%M:%S")
				now = datetime.strptime(now,"%Y-%m-%d %H:%M:%S")
				if datetime.strptime(start,"%Y-%m-%d %H:%M:%S") < now and datetime.strptime(end,"%Y-%m-%d %H:%M:%S") > now:
					results = StudentTI.query.filter_by(name=name,test_id=test_id).first()
					if results is not None:
						is_completed = results.completed
						if not is_completed:
							time_left = results.time_left
							if time_left <= duration:
								duration = time_left
								results = Student.query.filter_by(name=name,test_id=test_id).all()
								marked_ans = {}
								if results is not None:
									for row in results:
										marked_ans[row.question_id] = row.ans
									marked_ans = json.dumps(marked_ans)
									#return redirect(url_for('test' , testid = test_id))
						else:
							flash('Test already given', 'success')
							return redirect(url_for('give_test',val=val))
					else:
						info = StudentTI(name=name,test_id=test_id,time_left=duration)	
						db.session.add(info)
						db.session.commit()
						results = StudentTI.query.filter_by(name=name,test_id=test_id).first()
						if results is not None:
							is_completed = results.completed
							if not is_completed:
								time_left = results.time_left
								if time_left <= duration:
									duration = time_left
									results = Student.query.filter_by(name=name,test_id=test_id).all()
									marked_ans = {}
									if results is not None:
										for row in results:
											marked_ans[row.question_id] = row.ans
										marked_ans = json.dumps(marked_ans)
				else:
					if datetime.strptime(start,"%Y-%m-%d %H:%M:%S") > now:
						flash(f'Test start time is {start}', 'danger')
					else:
						flash(f'Test has ended', 'danger')
					return redirect(url_for('show_test'))
				return redirect(url_for('test' , testid = test_id))
			else:
				flash('Invalid password', 'danger')
				return redirect(url_for('give_test',val=val))
		flash('Invalid testid', 'danger')
		return redirect(url_for('give_test',val=val))
	return render_template('give_test.html', form = form,data=data1)

@app.route('/give-test/<testid>', methods=['GET','POST'])
@is_logged
def test(testid):
	global duration,marked_ans
	name = session['user']
	if request.method == 'GET':
		try:
			data = {'duration': duration, 'marks': '', 'q': '', 'a': "", 'b':"",'c':"",'d':"" }
			return render_template('quiz.html' ,**data, answers=marked_ans)
		except:
			return redirect(url_for('show_test'))
	else:
		flag = request.form['flag']

		if flag == 'get':
			num = request.form['no']
			results = T_question.query.filter_by(test_id=testid,question_id=num).first()
			if results is not None:
				data = results
				data = dict(results.__dict__); 
				data.pop('_sa_instance_state', None)
				return json.dumps(data)
		elif flag=='mark':
			qid = request.form['qid']
			ans = request.form['ans']
			results = Student.query.filter_by(name=name,test_id=testid,question_id=qid).first()
			if results is not None:
				results.ans = ans
			else:
				temp=Student(name=name,test_id=testid,question_id=qid,ans=ans)
				db.session.add(temp)
			db.session.commit()
		elif flag=='time':
			time_left = request.form['time']
			try:
				temp=StudentTI.query.filter_by(name=name,test_id=testid).first()
				temp.time_left=time_left
				db.session.commit()
			except:
				pass
		else:
			temp=StudentTI.query.filter_by(name=name,test_id=testid).first()
			temp.completed = True
			db.session.commit()
			flash("Test submitted successfully")
			return json.dumps({'sql':'fired'})

@app.route('/randomize', methods = ['POST'])
def random_gen():
	if request.method == "POST":
		id = request.form['id']
		name = session['user']
		results = T_question.query.filter_by(test_id=id).count()
		if results is not None:
			total = results
			nos = list(range(1,int(total)+1))
			print(nos)
			for no in nos:
				result = Student.query.filter_by(name=name,test_id=id,question_id=str(no)).first()
				if result is None:
					temp=Student(name=name,test_id=id,question_id=str(no),ans=' ')
					db.session.add(temp)	
			db.session.commit()
			random.Random(id).shuffle(nos)
			return json.dumps(nos)

def neg_marks(username,testid):
	results =  db2.execute("select marks,q.question_id as qid,q.c_ans as correct, s.ans as marked from tquestions q left join students s on  s.test_id = q.test_id and s.test_id = :testid and s.name = :username and s.question_id = q.question_id group by qid,correct,marked,marks order by q.question_id asc",{"testid":testid,"username":username}).fetchall()
	sum=0.0
	for i in results:
		if(str(i.marked) is not None):
			if(str(i.marked).upper() != str(i.correct)):
				sum=sum-0.25*int(i.marks)
			elif(str(i.marked).upper() == str(i.correct)):
				sum+=int(i.marks)
	return sum
def pos_marks(username,testid):
	results =  db2.execute("select marks,q.question_id as qid,q.c_ans as correct, s.ans as marked from tquestions q left join students s on  s.test_id = q.test_id and s.test_id = :testid and s.name = :username and s.question_id = q.question_id group by qid,correct,marked,marks order by q.question_id asc",{"testid":testid,"username":username}).fetchall()
	sum=0.0
	for i in results:
		if(str(i.marked) is not None):
			if(str(i.marked).upper() == str(i.correct)):
				sum+=int(i.marks)
	return sum	

def totmarks(username,tests): 
	temp = []
	for test in tests:
		testid = test['test_id']
		d = dict(test.items())
		results=db2.execute("select neg from test where test_id=:testid",{"testid":testid}).fetchone()
		if results['neg']:
			d['marks'] = neg_marks(username,testid) 

		else:
			d['marks'] = pos_marks(username,testid) 
		temp.append(d)
	print(temp)	
	return temp

@app.route('/<username>/tests-given')
@is_logged
def tests_given(username):
	if username == session['user']:
		results = db2.execute('select distinct(s_testinfo.test_id),subject,topic from s_testinfo,test where s_testinfo.name = :name and s_testinfo.test_id=test.test_id',{"name":username}).fetchall()
		results=totmarks(username,results)
		email = User.query.filter_by(name=username).first()
		return render_template('tests_given.html', tests=results,user=username,email=email.email,info=['home','logout'])
	else:
		flash('You are not authorized', 'danger')
		return redirect(url_for('home'))

@app.route('/<username>/<testid>')
@is_logged
def check_result(username, testid):
	if username == session['user']:
		results = db2.execute('SELECT * FROM test where test_id = :testid',{"testid":testid}).fetchone()
		if results is not None:
			check = results['show_result']
			if check:
				results = db2.execute('select explanation,question,a,b,c,d,marks,q.question_id as qid, \
					q.c_ans as correct, s.ans as marked from tquestions q inner join \
					students s on  q.test_id = s.test_id where q.question_id=s.question_id and q.test_id =:testid and s.name =:username \
					group by explanation,question,a,b,c,d,marks,qid,correct,marked\
					order by q.question_id asc', {'testid':testid,'username':username}).fetchall()
				
				if results is not None:
					return render_template('tests_result.html', results= results,i=0)
			else:
				flash('You are not authorized to check the result', 'danger')
				return redirect(url_for('tests_given',username = username))
	else:
		return redirect(url_for('home'))

@app.route('/<username>/tests-created')
@is_logged
@is_admin
def tests_created(username):
	if username == session['user']:
		results = db2.execute('select * from test where name = :username', {"username":username}).fetchall()
		return render_template('tests_created.html', tests=results,info=['home','logout'])
	else:
		flash('You are not authorized', 'danger')
		return redirect(url_for('home'))

def marks_calc(username,testid):
	results=db2.execute("select neg from test where test_id=:testid",{"testid":testid}).fetchone()
	if results['neg']:
		return neg_marks(username,testid) 
	return pos_marks(username,testid) 

@app.route('/<username>/tests-created/<testid>', methods = ['POST','GET'])
@is_logged
@is_admin
def student_results(username, testid):
	if username == session['user']:
		results = db2.execute('select user_details.name as name,test_id from s_testinfo,user_details where test_id = :testid and completed =true and s_testinfo.name=user_details.name ',{"testid":testid}).fetchall()
		final = []
		temp = []
		count = 1
		for user in results:
			score = marks_calc(user['name'], testid)
			d=dict(user.items())
			d['srno'] = count
			d['marks'] = score
			temp.append(d)
			final.append([count, user['name'], score])
			count+=1
		if request.method =='GET':
			results = temp
			return render_template('student_results.html', data=results)
		else:
			print(final)
			fields = ['Sr No', 'Name', 'Marks']
			with open('static/' + testid + '.csv', 'w') as f:
				writer = csv.writer(f)
				writer.writerow(fields)
				for row in final:
					writer.writerow(row)
			return send_file('static/' + testid + '.csv', as_attachment=True)

@app.route('/<username>/tests-created/<testid>/questions', methods = ['POST','GET'])
@is_logged
@is_admin
def questions(username, testid):
	if username == session['user']:
		results = db2.execute('SELECT * FROM test where test_id = :testid',{"testid":testid}).fetchall()
		if results is not None:
			results = db2.execute('select explanation,question,a,b,c,d,marks,q.question_id as qid, \
					q.c_ans as correct, s.ans as marked from tquestions q left join \
					students s on  s.test_id = q.test_id and s.test_id =:testid \
					and s.name =:username and s.question_id = q.question_id \
					group by explanation,question,a,b,c,d,marks,qid,correct,marked\
					order by q.question_id asc', {'testid':testid,'username':username}).fetchall()
			if results is not None:
				return render_template('disp_questions.html', results= results)